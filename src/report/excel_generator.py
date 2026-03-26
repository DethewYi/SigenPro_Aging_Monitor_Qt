"""CSV and Excel report generators for aging test records."""

import csv
import json
import logging
import os
from datetime import datetime, timedelta

logger = logging.getLogger(__name__)

# Result integer values (matching TestResult enum ordinals)
_RESULT_MAP = {1: "PASSED", 2: "FAILED", 3: "INTERRUPTED", 0: "PENDING"}


class ExcelReportGenerator:
    """Generate CSV (and optionally Excel) aging test reports."""

    def __init__(self):
        pass

    def _get_db(self):
        from ..storage.database_manager import DatabaseManager
        return DatabaseManager.instance()

    # ------------------------------------------------------------------
    # CSV report
    # ------------------------------------------------------------------

    def generate_report(self, test_record_id: int, file_path: str) -> bool:
        """Generate a UTF-8 BOM CSV report for *test_record_id*.

        Sections:
            Summary metadata, statistics, alarms, raw data.

        Returns:
            True on success, False on failure.
        """
        try:
            db = self._get_db()
            local = db.local_db

            # --- Load test record ---
            record = self._load_test_record(local, test_record_id)
            if record is None:
                logger.error("Test record %d not found", test_record_id)
                return False

            # Load template name
            template_name = self._load_template_name(local, record.get("template_id"))

            start_time = record.get("start_time", "")
            end_time = record.get("end_time") or datetime.now().isoformat()

            # --- Load alarms ---
            alarms = self._load_alarms(local, start_time, end_time)

            # --- Calculate statistics ---
            stats = self._calculate_statistics(local, start_time, end_time)

            # --- Load raw device data ---
            raw_data = self._load_raw_data(local, start_time, end_time)

            # --- Write CSV with BOM ---
            with open(file_path, "w", encoding="utf-8-sig", newline="") as f:
                writer = csv.writer(f)

                # Section 1: Summary
                result_str = _RESULT_MAP.get(record.get("result", 0), "UNKNOWN")
                writer.writerow(["== Summary =="])
                writer.writerow(["Record ID", record.get("record_id", "")])
                writer.writerow(["SN", record.get("device_sn", "")])
                writer.writerow(["PN", record.get("device_pn", "")])
                writer.writerow(["Template", template_name])
                writer.writerow(["Channel", record.get("channel_id", "")])
                writer.writerow(["Start Time", record.get("start_time", "")])
                writer.writerow(["End Time", record.get("end_time", "")])
                writer.writerow(["Result", result_str])
                writer.writerow(["Remark", record.get("remark", "")])
                writer.writerow([])

                # Section 2: Statistics
                writer.writerow(["== Statistics =="])
                writer.writerow(["Parameter", "Min", "Max", "Avg"])
                for pname, s in stats.items():
                    writer.writerow([
                        pname,
                        f"{s['min']:.4f}",
                        f"{s['max']:.4f}",
                        f"{s['avg']:.4f}",
                    ])
                writer.writerow([])

                # Section 3: Alarms
                writer.writerow(["== Alarm Records =="])
                writer.writerow(["Timestamp", "Device", "Parameter", "Value", "Threshold", "Direction"])
                for a in alarms:
                    direction = ">" if a.get("is_upper_limit") else "<"
                    writer.writerow([
                        a.get("timestamp", ""),
                        a.get("device_name", ""),
                        a.get("param_name", ""),
                        f"{a.get('current_value', 0):.4f}",
                        f"{a.get('threshold', 0):.4f}",
                        direction,
                    ])
                writer.writerow([])

                # Section 4: Raw Data
                writer.writerow(["== Raw Data =="])
                param_names = sorted(set(
                    k for row in raw_data for k in row.get("parameters", {}).keys()
                ))
                writer.writerow(["Timestamp"] + param_names)
                for rd in raw_data:
                    params = rd.get("parameters", {})
                    row_data = [rd.get("timestamp", "")]
                    for pname in param_names:
                        row_data.append(params.get(pname, ""))
                    writer.writerow(row_data)

            logger.info("CSV report generated: %s", file_path)
            return True

        except Exception:
            logger.exception("Failed to generate CSV report for record %d", test_record_id)
            return False

    # ------------------------------------------------------------------
    # Excel report (optional, with openpyxl fallback to CSV)
    # ------------------------------------------------------------------

    def generate_excel_report(self, test_record_id: int, file_path: str) -> bool:
        """Generate an Excel (.xlsx) report using openpyxl.

        If openpyxl is not available, falls back to CSV with the same file_path
        (changing extension to .csv).

        Returns:
            True on success, False on failure.
        """
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            logger.info("openpyxl not available, falling back to CSV")
            csv_path = os.path.splitext(file_path)[0] + ".csv"
            return self.generate_report(test_record_id, csv_path)

        try:
            db = self._get_db()
            local = db.local_db

            # --- Load data ---
            record = self._load_test_record(local, test_record_id)
            if record is None:
                logger.error("Test record %d not found", test_record_id)
                return False

            template_name = self._load_template_name(local, record.get("template_id"))
            start_time = record.get("start_time", "")
            end_time = record.get("end_time") or datetime.now().isoformat()
            alarms = self._load_alarms(local, start_time, end_time)
            stats = self._calculate_statistics(local, start_time, end_time)
            raw_data = self._load_raw_data(local, start_time, end_time)

            wb = openpyxl.Workbook()

            # --- Styles ---
            header_font = Font(bold=True, size=12)
            title_font = Font(bold=True, size=14)
            header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2",
                                      fill_type="solid")
            thin_border = Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin"),
            )

            # --- Sheet 1: Summary ---
            ws_summary = wb.active
            ws_summary.title = "\u6c47\u603b"  # "Summary" in Chinese

            result_str = _RESULT_MAP.get(record.get("result", 0), "UNKNOWN")
            summary_data = [
                ("Record ID", record.get("record_id", "")),
                ("SN", record.get("device_sn", "")),
                ("PN", record.get("device_pn", "")),
                ("Template", template_name),
                ("Channel", record.get("channel_id", "")),
                ("Start Time", record.get("start_time", "")),
                ("End Time", record.get("end_time", "")),
                ("Result", result_str),
                ("Remark", record.get("remark", "")),
            ]
            ws_summary.cell(row=1, column=1, value="Test Summary").font = title_font
            for i, (key, val) in enumerate(summary_data, start=3):
                cell_k = ws_summary.cell(row=i, column=1, value=key)
                cell_k.font = Font(bold=True)
                cell_v = ws_summary.cell(row=i, column=2, value=val)

            # Statistics section in the same sheet
            stat_start_row = 3 + len(summary_data) + 2
            ws_summary.cell(
                row=stat_start_row, column=1, value="Statistics"
            ).font = title_font

            stat_headers = ["Parameter", "Min", "Max", "Avg"]
            for j, h in enumerate(stat_headers, start=1):
                cell = ws_summary.cell(row=stat_start_row + 1, column=j, value=h)
                cell.font = header_font
                cell.fill = header_fill
                cell.border = thin_border

            for i, (pname, s) in enumerate(stats.items(), start=stat_start_row + 2):
                ws_summary.cell(row=i, column=1, value=pname).border = thin_border
                ws_summary.cell(row=i, column=2, value=round(s["min"], 4)).border = thin_border
                ws_summary.cell(row=i, column=3, value=round(s["max"], 4)).border = thin_border
                ws_summary.cell(row=i, column=4, value=round(s["avg"], 4)).border = thin_border

            # Auto-width for summary
            ws_summary.column_dimensions["A"].width = 18
            ws_summary.column_dimensions["B"].width = 20
            ws_summary.column_dimensions["C"].width = 14
            ws_summary.column_dimensions["D"].width = 14

            # --- Sheet 2: Alarms ---
            ws_alarms = wb.create_sheet(title="\u62a5\u8b66")  # "Alarms"
            alarm_headers = ["Timestamp", "Device", "Parameter", "Value", "Threshold", "Direction"]
            for j, h in enumerate(alarm_headers, start=1):
                cell = ws_alarms.cell(row=1, column=j, value=h)
                cell.font = header_font
                cell.fill = header_fill
                cell.border = thin_border

            for i, a in enumerate(alarms, start=2):
                direction = ">" if a.get("is_upper_limit") else "<"
                ws_alarms.cell(row=i, column=1, value=a.get("timestamp", "")).border = thin_border
                ws_alarms.cell(row=i, column=2, value=a.get("device_name", "")).border = thin_border
                ws_alarms.cell(row=i, column=3, value=a.get("param_name", "")).border = thin_border
                ws_alarms.cell(row=i, column=4, value=a.get("current_value", 0)).border = thin_border
                ws_alarms.cell(row=i, column=5, value=a.get("threshold", 0)).border = thin_border
                ws_alarms.cell(row=i, column=6, value=direction).border = thin_border

            # --- Sheet 3: Raw Data ---
            ws_raw = wb.create_sheet(title="\u539f\u59cb\u6570\u636e")  # "Raw Data"
            param_names = sorted(set(
                k for row in raw_data for k in row.get("parameters", {}).keys()
            ))
            raw_headers = ["Timestamp"] + param_names
            for j, h in enumerate(raw_headers, start=1):
                cell = ws_raw.cell(row=1, column=j, value=h)
                cell.font = header_font
                cell.fill = header_fill
                cell.border = thin_border

            for i, rd in enumerate(raw_data, start=2):
                params = rd.get("parameters", {})
                ws_raw.cell(row=i, column=1, value=rd.get("timestamp", "")).border = thin_border
                for j, pname in enumerate(param_names, start=2):
                    ws_raw.cell(row=i, column=j, value=params.get(pname, "")).border = thin_border

            # Auto-width for raw data sheet
            ws_raw.column_dimensions["A"].width = 22
            for j in range(2, len(raw_headers) + 1):
                col_letter = openpyxl.utils.get_column_letter(j)
                ws_raw.column_dimensions[col_letter].width = 14

            wb.save(file_path)
            logger.info("Excel report generated: %s", file_path)
            return True

        except Exception:
            logger.exception("Failed to generate Excel report for record %d", test_record_id)
            # Fallback to CSV
            csv_path = os.path.splitext(file_path)[0] + ".csv"
            return self.generate_report(test_record_id, csv_path)

    # ------------------------------------------------------------------
    # Shared data loading helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _load_test_record(local_db, test_record_id: int) -> dict | None:
        row = local_db._conn.execute(
            "SELECT record_id, template_id, channel_id, device_sn, device_pn, "
            "start_time, end_time, result, remark FROM test_records WHERE record_id=?",
            (test_record_id,),
        ).fetchone()
        if row is None:
            return None
        return {
            "record_id": row[0],
            "template_id": row[1],
            "channel_id": row[2],
            "device_sn": row[3],
            "device_pn": row[4],
            "start_time": row[5],
            "end_time": row[6],
            "result": row[7],
            "remark": row[8],
        }

    @staticmethod
    def _load_template_name(local_db, template_id: int | None) -> str:
        if not template_id or template_id <= 0:
            return ""
        tmpl = local_db.load_template(template_id)
        return tmpl.get("name", "") if tmpl else ""

    @staticmethod
    def _load_alarms(local_db, start_time: str, end_time: str) -> list[dict]:
        rows = local_db._conn.execute(
            "SELECT timestamp, device_name, param_name, current_value, "
            "threshold, is_upper_limit FROM alarms "
            "WHERE timestamp >= ? AND timestamp <= ? ORDER BY timestamp",
            (start_time, end_time),
        ).fetchall()
        return [
            {
                "timestamp": r[0],
                "device_name": r[1],
                "param_name": r[2],
                "current_value": r[3],
                "threshold": r[4],
                "is_upper_limit": bool(r[5]),
            }
            for r in rows
        ]

    @staticmethod
    def _load_raw_data(local_db, start_time: str, end_time: str) -> list[dict]:
        """Load all device data rows from daily tables in the time range."""
        try:
            start_dt = datetime.fromisoformat(start_time)
        except (ValueError, TypeError):
            start_dt = datetime.now()
        try:
            end_dt = datetime.fromisoformat(end_time)
        except (ValueError, TypeError):
            end_dt = datetime.now()

        start_iso = start_dt.isoformat()
        end_iso = end_dt.isoformat()
        results = []
        current = start_dt

        while current <= end_dt:
            date_str = current.strftime("%Y%m%d")
            table_name = f"device_data_{date_str}"

            count = local_db._conn.execute(
                "SELECT count(*) FROM sqlite_master WHERE type='table' AND name=?",
                (table_name,),
            ).fetchone()[0]

            if count > 0:
                rows = local_db._conn.execute(
                    f"SELECT timestamp, parameters FROM [{table_name}] "
                    f"WHERE timestamp>=? AND timestamp<=? ORDER BY timestamp",
                    (start_iso, end_iso),
                ).fetchall()
                for r in rows:
                    try:
                        params = json.loads(r[1])
                    except (json.JSONDecodeError, TypeError):
                        params = {}
                    results.append({"timestamp": r[0], "parameters": params})

            current += timedelta(days=1)

        return results

    @staticmethod
    def _calculate_statistics(local_db, start_time: str, end_time: str) -> dict:
        """Calculate min/max/avg per parameter across all device data."""
        raw_data = ExcelReportGenerator._load_raw_data(local_db, start_time, end_time)

        param_values: dict[str, list[float]] = {}
        for row in raw_data:
            params = row.get("parameters", {})
            for pname, pval in params.items():
                try:
                    val = float(pval)
                    param_values.setdefault(pname, []).append(val)
                except (ValueError, TypeError):
                    continue

        stats: dict[str, dict] = {}
        for pname, values in param_values.items():
            if values:
                stats[pname] = {
                    "min": min(values),
                    "max": max(values),
                    "avg": sum(values) / len(values),
                }
        return stats
